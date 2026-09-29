import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.contrib import admin
from django.http import HttpResponse
from .models import SolicitudTerreno, Docente, Asignatura, UnidadAprendizaje, AprendizajeEsperado


@admin.action(description='Descargar seleccionadas como Excel (.xlsx)')
def exportar_solicitudes_excel(modeladmin, request, queryset):
    """
    Acción de administración para exportar solicitudes seleccionadas a Excel.
    """
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="solicitudes_terreno.xlsx"'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Solicitudes"

    headers = [
        'ID', 
        'Docente', 
        'Asignatura', 
        'Unidad de Aprendizaje', 
        'Aprendizaje Esperado', 
        'Fecha Propuesta', 
        'Duración (Horas)',
        'Estado'
    ]
    ws.append(headers)

    # Estilo de cabecera institucional
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    solicitudes = queryset.select_related('docente', 'asignatura', 'unidad', 'aprendizaje_esperado')

    for solicitud in solicitudes:
        fecha_str = ''
        if hasattr(solicitud, 'fecha_propuesta') and solicitud.fecha_propuesta:
            fecha_str = solicitud.fecha_propuesta.strftime('%d/%m/%Y')

        row = [
            solicitud.id,
            str(solicitud.docente) if solicitud.docente else 'N/A',
            str(solicitud.asignatura) if solicitud.asignatura else 'N/A',
            str(solicitud.unidad) if solicitud.unidad else 'N/A',
            str(solicitud.aprendizaje_esperado) if solicitud.aprendizaje_esperado else 'N/A',
            fecha_str,
            getattr(solicitud, 'duracion_horas', 'N/A'),
            solicitud.get_estado_display() if hasattr(solicitud, 'get_estado_display') else solicitud.estado
        ]
        ws.append(row)

        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=ws.max_row, column=col_num)
            cell.border = thin_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(response)
    return response


@admin.register(Docente)
class DocenteAdmin(admin.ModelAdmin):
    list_display = ('id', 'rut', 'nombre', 'email')
    search_fields = ('rut', 'nombre', 'email')


@admin.register(Asignatura)
class AsignaturaAdmin(admin.ModelAdmin):
    list_display = ('id', 'codigo', 'nombre', 'seccion', 'get_docentes')
    list_filter = ('docentes',)
    search_fields = ('codigo', 'nombre', 'docentes__nombre', 'docentes__rut')

    @admin.display(description='Docentes')
    def get_docentes(self, obj):
        """Muestra los nombres de todos los docentes asignados en la tabla del panel."""
        return ", ".join([d.nombre for d in obj.docentes.all()]) if obj.docentes.exists() else "Sin docente"


@admin.register(UnidadAprendizaje)
class UnidadAprendizajeAdmin(admin.ModelAdmin):
    list_display = ('id', 'numero', 'nombre', 'asignatura')
    list_filter = ('asignatura',)
    search_fields = ('nombre', 'asignatura__nombre', 'asignatura__codigo')


@admin.register(AprendizajeEsperado)
class AprendizajeEsperadoAdmin(admin.ModelAdmin):
    list_display = ('id', 'codigo', 'descripcion', 'unidad')
    list_filter = ('unidad__asignatura',)
    search_fields = ('codigo', 'descripcion')


@admin.register(SolicitudTerreno)
class SolicitudTerrenoAdmin(admin.ModelAdmin):
    list_display = ('id', 'docente', 'asignatura', 'unidad', 'fecha_propuesta', 'estado', 'contratada', 'realizada')
    list_filter = ('estado', 'contratada', 'realizada', 'docente', 'asignatura')
    search_fields = ('docente__nombre', 'asignatura__nombre', 'institucion_destino')
    actions = [exportar_solicitudes_excel]